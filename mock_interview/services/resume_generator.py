"""Generate an ATS-friendly resume as editable DOCX and printable PDF.

The download is a single-column, machine-readable rebuild of the uploaded
resume.  It follows the reference hierarchy:

    Header -> Objective -> Education -> Skills -> Real-Time Projects ->
    Certifications -> Internships & Experience -> Academic Projects ->
    Achievements

Tailoring is evidence-based.  The target role is supplied by the student (or
derived from the job description); existing skills are grouped into labelled
categories; projects that the upload labels as real-time or academic stay in
separate sections; and job-description terms are used only to order existing
evidence and skills, never to add unsupported qualifications.  An optional
local-LLM wording pass is accepted only after the evidence checks in
``resume_tailoring.filter_rewrite_map`` pass.

"Personal particulars" / "References" blocks (date of birth, gender, IDs, ...)
are excluded because they hurt ATS parsing, exactly as the checker advises.
"""

import io
import re

from django.core.files.base import ContentFile

from mock_interview.services.ats import (
    _BULLET_RE,
    _EMAIL_RE,
    _GITHUB_RE,
    _LINKEDIN_RE,
    _PORTFOLIO_RE,
    _SECTION_SYNONYMS,
    _detect_sections,
    _lines,
    _normalize_heading,
)
from mock_interview.services.resume_tailoring import (
    _SKILL_CATEGORY_BY_KEY,
    build_objective,
    detect_career_level,
    filter_rewrite_map,
    group_skills,
    infer_target_role,
)

# Reference order from the supplied fresher-resume structure.  Sections that
# have no evidence are omitted rather than emitted empty.
SECTION_ORDER = (
    ("summary", "OBJECTIVE"),
    ("education", "EDUCATION"),
    ("skills", "SKILLS"),
    ("real_projects", "REAL-TIME PROJECTS"),
    ("projects", "PROJECTS"),
    ("certifications", "CERTIFICATIONS"),
    ("experience", "INTERNSHIPS & EXPERIENCE"),
    ("academic_projects", "ACADEMIC PROJECTS"),
    ("achievements", "ACHIEVEMENTS"),
)

EXTRA_SECTIONS = (
    ("activities", "ACTIVITIES"),
    ("coursework", "COURSEWORK"),
    ("publications", "PUBLICATIONS"),
    ("languages", "LANGUAGES"),
    ("interests", "INTERESTS"),
)

# Canonical sections deliberately excluded from the ATS-friendly rebuild.
_SKIPPED_SECTIONS = frozenset(("personal", "references"))

_OBJECTIVE_HEADINGS = frozenset(
    (
        "objective", "career objective", "professional objective",
        "job objective", "objective statement",
    )
)

_SECTION_HEADINGS = frozenset(
    name
    for names in _SECTION_SYNONYMS.values()
    for name in names
)
_CONTACT_HEADINGS = frozenset(_SECTION_SYNONYMS["contact"])
_NAME_STOPWORDS = frozenset({
    "resume", "curriculum vitae", "cv", "bio-data", "biodata", "name",
})
_CITY_SUFFIX_RE = re.compile(
    r"\b([A-Z][A-Za-z.'-]+(?:\s+[A-Z][A-Za-z.'-]+){0,2}\s+City)\b"
)
_STREET_RE = re.compile(
    r"\b(?:street|st\.?|avenue|ave\.?|road|rd\.?|lane|ln\.?|drive|dr\.?|"
    r"boulevard|blvd\.?|barangay|brgy\.?)\b",
    re.IGNORECASE,
)
_COUNTRY_ABBR = frozenset(
    {"us", "u.s.", "usa", "u.s.a.", "uk", "u.k.", "uae", "u.a.e."}
)
_COUNTRY_NAMES = _COUNTRY_ABBR | frozenset({
    "philippines", "india", "united states", "united states of america",
    "united kingdom", "united arab emirates", "canada", "australia",
    "singapore", "malaysia", "new zealand", "ireland", "germany", "france",
    "netherlands", "japan", "south korea", "china", "south africa",
    "pakistan", "bangladesh", "sri lanka", "nepal", "saudi arabia", "qatar",
    "oman", "kuwait",
})
_PHONE_CANDIDATE_RE = re.compile(
    r"(?<![\w])(?:\+?\d{1,3}[\s.-]?)?(?:\(\d{2,5}\)[\s.-]?)?"
    r"\d{3,5}(?:[\s.-]?\d{3,5}){1,2}(?![\w])"
)
_DATE_RANGE_RE = re.compile(r"^(?:19|20)\d{2}\s*[-–—]\s*(?:19|20)\d{2}$")


# ── small helpers ─────────────────────────────────────────────────────────
def _profile(profile):
    return profile if isinstance(profile, dict) else {}


def _clean_lines(lines):
    """Trim, drop empties/separators and remove exact duplicates in order."""
    seen = set()
    out = []
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if re.fullmatch(r"[-=_*#~•·…]{3,}", line):
            continue
        key = line.lower()
        if key in seen:
            continue
        seen.add(key)
        out.append(line)
    return out


def _strip_marker(line):
    removed = _BULLET_RE.sub("", line).strip()
    return removed or line


def _item(line, is_bullet=None):
    marked = bool(_BULLET_RE.match(line))
    kind = "bullet" if (is_bullet if is_bullet is not None else marked) else "plain"
    return {"text": _strip_marker(line), "kind": kind}


def _profile_lines(profile, key, source_text=""):
    """Profile fallback lines, restricted to text actually present in the upload.

    The structured profile is model-generated, so it is never trusted on its
    own.  A line is used only when it can be found in the extracted resume.
    This prevents a stale or hallucinated profile value from entering the
    downloadable file.
    """
    source = re.sub(r"\s+", " ", str(source_text or "")).strip().lower()
    out = []
    for item in profile.get(key) or []:
        for line in _lines(str(item)):
            if source and re.sub(r"\s+", " ", line).strip().lower() not in source:
                continue
            out.append(line)
    return out


def _header_lines(lines):
    """Return the header, including a labelled CONTACT block.

    A contact heading is part of the header rather than the end of it; any
    other section heading marks the end.
    """
    header = []
    for line in lines[:12]:
        norm = _normalize_heading(line)
        if norm in _SECTION_HEADINGS:
            if norm in _CONTACT_HEADINGS:
                continue
            break
        header.append(line)
    return header


def _find_phone(text):
    """Return the first complete-looking phone number, not a date range."""
    for match in _PHONE_CANDIDATE_RE.finditer(str(text or "")):
        candidate = match.group(0).strip().strip(".,;|")
        if not candidate or _DATE_RANGE_RE.match(candidate):
            continue
        digits = re.sub(r"\D", "", candidate)
        if len(digits) < 7:
            continue
        return candidate
    return ""


def _detect_name(resume, lines):
    """Prefer the name in the upload, not an account ID used as a username."""
    account_name = re.sub(
        r"\s+", " ", str(getattr(resume, "student_name", "") or "")
    ).strip()
    header = _header_lines(lines)
    account_low = account_name.lower()
    placeholder_account = (
        not account_name
        or bool(re.search(r"\d", account_name))
        or bool(re.search(r"[@/]", account_name))
        or account_low in {
            "student", "candidate", "user", "admin", "resume", "your name",
        }
    )
    if account_name and not placeholder_account:
        for line in header:
            if re.sub(r"\s+", " ", line).strip().lower() == account_low:
                return line.strip().rstrip(".,;")

    for line in header[:6]:
        if _normalize_heading(line) in _SECTION_HEADINGS:
            continue
        candidate = re.split(
            r"\s*[|•·]\s*|\s+[-–—]\s+", line, maxsplit=1
        )[0].strip()
        candidate = re.sub(
            r"^(?:name|full name)\s*[:\-]\s*",
            "",
            candidate,
            flags=re.IGNORECASE,
        ).strip()
        if not candidate:
            continue
        if candidate.lower().strip(" .:") in _NAME_STOPWORDS:
            continue
        if _EMAIL_RE.search(candidate) or _find_phone(candidate):
            continue
        if any(
            pattern.search(candidate)
            for pattern in (_LINKEDIN_RE, _GITHUB_RE, _PORTFOLIO_RE)
        ) or "://" in candidate:
            continue
        words = candidate.split()
        if 1 <= len(words) <= 5 and candidate[:1].isalpha():
            return candidate.rstrip(".,;") or "Your Name"
    if account_name and not placeholder_account:
        return account_name
    return "Your Name"


def _detect_location(lines, name=""):
    # Only inspect the header area and stop at a section heading.  Without this
    # guard a skills line such as "Python, SQL" can look like "City, State".
    name_low = re.sub(r"\s+", " ", name or "").strip().lower()
    for line in _header_lines(lines):
        if _EMAIL_RE.search(line) or _find_phone(line):
            continue
        for segment in re.split(r"\s*[|•·]\s*", line):
            segment = re.sub(
                r"^(?:address|location|city|based in)\s*[:\-]\s*",
                "",
                segment,
                flags=re.IGNORECASE,
            ).strip()
            if not segment:
                continue
            if name_low and segment.lower().strip(" .,;") == name_low:
                continue

            city = _CITY_SUFFIX_RE.search(segment)
            city_name = city.group(1).strip(" ,") if city else ""

            parts = [
                part.strip() for part in segment.split(",") if part.strip()
            ]
            if len(parts) < 2:
                if city_name:
                    return city_name
                continue
            first, last = parts[0], parts[-1]
            if first.lower() in _SKILL_CATEGORY_BY_KEY:
                continue
            if re.match(r"^\d", first) or _STREET_RE.search(first):
                return ", ".join(parts[1:]) if len(parts) > 1 else ""
            if re.fullmatch(r"[A-Z]{2}", last):
                return ", ".join(parts)
            if last.lower() in _COUNTRY_NAMES:
                return ", ".join(parts)
            if (
                len(parts) == 2
                and 1 <= len(first.split()) <= 3
                and re.fullmatch(r"[A-Z][A-Za-z.'-]+", last)
                and not last.isupper()
            ):
                return ", ".join(parts)
            if city_name:
                return city_name
    return ""


def _link_value(match):
    if not match:
        return ""
    return match.group(0).strip().rstrip(".,;:)|]")


def _build_header_from(name, text, lines):
    """Build the header line: location | phone | email | LinkedIn | links."""
    head = "\n".join(_header_lines(lines))
    fields = []
    location = _detect_location(lines, name=name)
    if location:
        fields.append(location)
    phone = _find_phone(head)
    if phone:
        fields.append(phone)
    email = _EMAIL_RE.search(text)
    if email:
        fields.append(email.group(0).strip())

    links = []
    for pattern in (_LINKEDIN_RE, _GITHUB_RE, _PORTFOLIO_RE):
        value = _link_value(pattern.search(text))
        if not value:
            continue
        if any(value.lower() in existing.lower()
               or existing.lower() in value.lower() for existing in links):
            continue
        links.append(value)
    fields.extend(links)
    return name, "  |  ".join(fields)


def build_header(resume, lines):
    """Return (name, contact_line) matching the sample layout's header block."""
    text = str(getattr(resume, "extracted_text", "") or "")
    return _build_header_from(_detect_name(resume, lines), text, lines)


def extract_contact(text):
    """Backward-compatible summarised contact line (phone | email | link)."""
    lines = _lines(text)
    _, contact = _build_header_from("", text, lines)
    return contact


_SENSITIVE_SUMMARY_RE = re.compile(
    r"\b(date of birth|d\.?o\.?b|passport|aadhaar|aadhar|marital status|"
    r"gender|father'?s name|mother'?s name|nationality|religion|"
    r"\bage\b\s*[:\-]?\s*\d)\b",
    re.IGNORECASE,
)


_SUMMARY_SIGNAL_RE = re.compile(
    r"\b(?:fresher|student|graduate|professional|engineer|developer|"
    r"specialist|seeking|objective|experienced|skilled|background|"
    r"aspiring|pursuing|final[- ]year|b\.?tech|b\.?e\.?|bachelor|master|"
    r"mca|bca|degree)\b",
    re.IGNORECASE,
)


def first_summary(text):
    """First summary-like paragraph, used only as a conservative fallback.

    Skips contact, bullet, and personal-information lines so the generated
    resume never surfaces sensitive details or treats a project bullet as an
    objective statement.
    """
    lines = [
        line.strip()
        for line in (text or "").splitlines()
        if line.strip()
    ]
    for line in lines[:10]:
        if len(line) < 60:
            continue
        if _BULLET_RE.match(line):
            continue
        if _EMAIL_RE.search(line) or _find_phone(line):
            continue
        if _SENSITIVE_SUMMARY_RE.search(line):
            continue
        if not _SUMMARY_SIGNAL_RE.search(line):
            continue
        return line[:600]
    return ""


def _fallback_summary(text):
    summary = first_summary(text)
    return [summary] if summary else []


def _has_objective_heading(lines):
    for line in lines:
        if len(line) > 46:
            continue
        norm = _normalize_heading(line)
        if norm in _OBJECTIVE_HEADINGS:
            return True
    return False


def _make_section(key, title, lines, *, bullet=True):
    content = _clean_lines(lines)
    if not content:
        return None
    return {
        "key": key,
        "title": title,
        "items": [_item(line, is_bullet=bullet) for line in content],
    }


def _dedupe_project_items(sections):
    """Remove exact duplicates across the project/achievement blocks."""
    keys = {"real_projects", "projects", "academic_projects", "achievements"}
    seen = set()
    for section in sections:
        if section.get("key") not in keys:
            continue
        kept = []
        for item in section.get("items", []):
            label = _item_text(item).strip().lower()
            if not label or label in seen:
                continue
            seen.add(label)
            kept.append(item)
        section["items"] = kept
    return [
        section for section in sections
        if section.get("is_contact") or section.get("items")
    ]


def _missing_detail_warnings(text, name, education, skills, projects, experience):
    warnings = []
    if not name or name.strip().lower() in {"your name", "resume"}:
        warnings.append(
            "No candidate name was detected; the header uses an editable "
            "placeholder instead of guessing."
        )
    if not _EMAIL_RE.search(text):
        warnings.append(
            "No email address was detected in the upload; none was invented."
        )
    header_text = "\n".join(_header_lines(_lines(text)))
    if not _find_phone(header_text):
        warnings.append(
            "No phone number was detected in the upload; none was invented."
        )
    if not education:
        warnings.append(
            "No education details were detected; the section was omitted "
            "rather than fabricated."
        )
    if not skills:
        warnings.append(
            "No skills section was detected; no skills were added from the "
            "job description."
        )
    if not projects and not experience:
        warnings.append(
            "No projects or experience were detected; those sections were "
            "omitted rather than filled with invented examples."
        )
    return warnings


def build_sections(
    resume,
    rewrite_map=None,
    *,
    target_role="",
    career_level="auto",
    job_description="",
    matched_keywords=None,
    warnings=None,
    meta=None,
):
    """Return ordered [{key, title, items}] sections for the ATS layout.

    ``rewrite_map`` is optional.  It is filtered by
    ``resume_tailoring.filter_rewrite_map`` before use, so a model rewrite can
    never add an unsupported technology, number, or named detail.
    """
    text = str(getattr(resume, "extracted_text", "") or "")
    profile = _profile(getattr(resume, "structured_profile", {}))
    lines = _lines(text)
    _present, spans = _detect_sections(lines)
    warnings = warnings if warnings is not None else []
    meta = meta if meta is not None else {}

    resolved_role = infer_target_role(
        job_description, target_role, text
    )
    resolved_level = detect_career_level(resume, career_level)
    meta["target_role"] = resolved_role
    meta["career_level"] = resolved_level

    name, contact = build_header(resume, lines)

    summary_lines = [
        line for line in spans.get("summary", [])
        if not _SENSITIVE_SUMMARY_RE.search(line)
    ]
    if not summary_lines:
        summary_lines = _fallback_summary(text)

    skills_lines = list(spans.get("skills", [])) or _profile_lines(
        profile, "skills", text
    )
    education = list(spans.get("education", [])) or _profile_lines(
        profile, "education", text
    )
    certifications = list(spans.get("certifications", [])) or _profile_lines(
        profile, "certifications", text
    )
    experience = (
        list(spans.get("experience", []))
        + list(spans.get("internships", []))
    )
    if not experience:
        experience = _profile_lines(profile, "experience", text)

    real_projects = list(spans.get("real_projects", []))
    projects = list(spans.get("projects", []))
    academic_projects = list(spans.get("academic_projects", []))
    if not (real_projects or projects or academic_projects):
        projects = _profile_lines(profile, "projects", text)

    achievements = list(spans.get("achievements", []))
    extras = {
        key: list(spans.get(key, []))
        for key, _title in EXTRA_SECTIONS
    }

    original = {
        "summary": summary_lines,
        "experience": experience,
        "projects": projects,
        "education": education,
        "certifications": certifications,
    }
    safe_rewrites = filter_rewrite_map(
        resume,
        original,
        rewrite_map,
        target_role=resolved_role,
        warnings=warnings,
    )
    if safe_rewrites:
        for key in ("experience", "projects", "education", "certifications"):
            if key in safe_rewrites:
                original[key] = safe_rewrites[key]
        experience = original["experience"]
        projects = original["projects"]
        education = original["education"]
        certifications = original["certifications"]
    meta["applied_rewrites"] = bool(safe_rewrites)
    meta["applied_keys"] = sorted(safe_rewrites)

    grouped_skills, skill_names = group_skills(
        skills_lines, matched_keywords=matched_keywords
    )

    objective = build_objective(
        summary_lines=summary_lines,
        rewritten_summary=safe_rewrites.get("summary"),
        skills=skill_names,
        has_projects=bool(real_projects or projects or academic_projects),
        has_experience=bool(experience),
        target_role=resolved_role,
        career_level=resolved_level,
    )

    sections = [{
        "key": "contact",
        "is_contact": True,
        "name": name,
        "contact": contact,
    }]

    summary_title = (
        "PROFESSIONAL SUMMARY"
        if resolved_level == "experienced"
        else "OBJECTIVE"
    )
    section = _make_section(
        "summary", summary_title, objective, bullet=False
    )
    if section:
        sections.append(section)

    section = _make_section(
        "education", "EDUCATION", education, bullet=False
    )
    if section:
        sections.append(section)

    section = _make_section("skills", "SKILLS", grouped_skills, bullet=True)
    if section:
        sections.append(section)

    section = _make_section(
        "real_projects", "REAL-TIME PROJECTS", real_projects
    )
    if section:
        sections.append(section)

    section = _make_section("projects", "PROJECTS", projects)
    if section:
        sections.append(section)

    section = _make_section(
        "certifications", "CERTIFICATIONS", certifications, bullet=True
    )
    if section:
        sections.append(section)

    section = _make_section(
        "experience", "INTERNSHIPS & EXPERIENCE", experience
    )
    if section:
        sections.append(section)

    section = _make_section(
        "academic_projects", "ACADEMIC PROJECTS", academic_projects
    )
    if section:
        sections.append(section)

    section = _make_section(
        "achievements", "ACHIEVEMENTS", achievements
    )
    if section:
        sections.append(section)

    # Remaining labelled sections are carried over so nothing is lost
    # (activities, coursework, publications, languages, interests).
    for key, title in EXTRA_SECTIONS:
        section = _make_section(key, title, extras.get(key, []))
        if section:
            sections.append(section)

    sections = _dedupe_project_items(sections)

    for message in _missing_detail_warnings(
        text, name, education, grouped_skills, real_projects or projects
        or academic_projects, experience
    ):
        if message not in warnings:
            warnings.append(message)
    for message in _quality_warnings(sections):
        if message not in warnings:
            warnings.append(message)
    return sections


def _item_text(item):
    return item["text"] if isinstance(item, dict) else str(item)


def _item_kind(item):
    return item.get("kind") if isinstance(item, dict) else "bullet"


# Conservative, resume-specific typo checks.  Correcting text automatically
# would risk changing names and technical terms, so these become warnings for
# the editable DOCX instead.
_COMMON_MISSPELLINGS = {
    "acheive": "achieve",
    "acheived": "achieved",
    "calender": "calendar",
    "cerification": "certification",
    "certifcate": "certificate",
    "commited": "committed",
    "comitted": "committed",
    "comunication": "communication",
    "definately": "definitely",
    "developement": "development",
    "enviroment": "environment",
    "enviornment": "environment",
    "expereince": "experience",
    "experiance": "experience",
    "intership": "internship",
    "maintainance": "maintenance",
    "maintanance": "maintenance",
    "managment": "management",
    "occured": "occurred",
    "occurance": "occurrence",
    "perfomance": "performance",
    "porgramming": "programming",
    "proffesional": "professional",
    "programing": "programming",
    "projet": "project",
    "recieve": "receive",
    "recieved": "received",
    "responsibilty": "responsibility",
    "seperate": "separate",
    "seperated": "separated",
    "sucess": "success",
    "sucessful": "successful",
    "technolgy": "technology",
    "technolgies": "technologies",
}
_REPEATED_WORD_RE = re.compile(r"\b([A-Za-z]{2,})\s+\1\b", re.IGNORECASE)


def _quality_warnings(sections):
    """Flag likely spelling and repeated-word issues without rewriting them."""
    warnings = []
    repeated = set()
    words = []
    for section in sections:
        for item in section.get("items", []):
            text = _item_text(item)
            words.extend(re.findall(r"[A-Za-z]+", text))
            for match in _REPEATED_WORD_RE.finditer(text):
                repeated.add(match.group(0))
    if repeated:
        shown = ", ".join(f"'{value}'" for value in sorted(repeated)[:5])
        warnings.append(
            f"Repeated word(s) detected in the upload: {shown}. Correct them "
            "in the editable DOCX before submitting."
        )
    misspelled = sorted({
        word for word in words if word.lower() in _COMMON_MISSPELLINGS
    }, key=str.lower)
    if misspelled:
        details = ", ".join(
            f"'{word}' (maybe '{_COMMON_MISSPELLINGS[word.lower()]}')"
            for word in misspelled[:8]
        )
        warnings.append(
            f"Possible spelling issue(s) detected: {details}. Correct them in "
            "the editable DOCX before submitting."
        )
    return warnings


def build_docx(sections):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    document = Document()
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    for section in sections:
        if section.get("is_contact"):
            name_paragraph = document.add_paragraph()
            name_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            name_run = name_paragraph.add_run(section.get("name") or "Resume")
            name_run.bold = True
            name_run.font.size = Pt(16)
            contact = section.get("contact") or ""
            if contact:
                contact_paragraph = document.add_paragraph()
                contact_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                contact_run = contact_paragraph.add_run(contact)
                contact_run.font.size = Pt(10)
                contact_run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            continue
        heading = document.add_heading(section["title"], level=2)
        for run in heading.runs:
            run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        for item in section.get("items", []):
            item_text = _item_text(item)
            if _item_kind(item) == "bullet":
                bullet = document.add_paragraph(style="List Bullet")
                bullet.add_run(item_text)
            else:
                paragraph = document.add_paragraph()
                paragraph.add_run(item_text)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def build_pdf(sections):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
    )

    def escape(value):
        return (
            str(value)
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )

    stream = io.BytesIO()
    doc = SimpleDocTemplate(
        stream,
        pagesize=A4,
        topMargin=0.6 * inch,
        bottomMargin=0.7 * inch,
        leftMargin=0.7 * inch,
        rightMargin=0.7 * inch,
    )
    styles = getSampleStyleSheet()
    name_style = ParagraphStyle(
        "ATSName", parent=styles["Normal"], fontName="Helvetica-Bold",
        fontSize=16, leading=19, alignment=TA_CENTER,
        textColor=colors.HexColor("#1a1a1a"), spaceAfter=1,
    )
    contact_style = ParagraphStyle(
        "ATSContact", parent=styles["Normal"], fontName="Helvetica",
        fontSize=9.5, leading=13, alignment=TA_CENTER,
        textColor=colors.HexColor("#333333"), spaceAfter=5,
    )
    heading_style = ParagraphStyle(
        "ATSHeading", parent=styles["Normal"], fontName="Helvetica-Bold",
        fontSize=11.5, leading=14, textColor=colors.HexColor("#1F3A5F"),
        spaceBefore=9, spaceAfter=3,
    )
    body_style = ParagraphStyle(
        "ATSBody", parent=styles["Normal"], fontName="Helvetica",
        fontSize=9.5, leading=13, spaceAfter=2,
    )
    bullet_style = ParagraphStyle(
        "ATSBullet", parent=body_style, leftIndent=14,
        firstLineIndent=-10, spaceAfter=1.5,
    )

    story = []
    for section in sections:
        if section.get("is_contact"):
            story.append(Paragraph(
                escape(section.get("name") or "Resume"), name_style
            ))
            if section.get("contact"):
                story.append(Paragraph(escape(section["contact"]), contact_style))
            continue
        story.append(Paragraph(escape(section["title"]), heading_style))
        for item in section.get("items", []):
            item_text = escape(_item_text(item))
            if _item_kind(item) == "bullet":
                # An ASCII hyphen stays legible and extractable in every PDF
                # parser; the bullet glyph can extract as a control character.
                story.append(Paragraph(f"-&nbsp;{item_text}", bullet_style))
            else:
                story.append(Paragraph(item_text, body_style))
    doc.build(story)
    return stream.getvalue()


class ResumeGenerationError(RuntimeError):
    pass


def _safe_filename(name):
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", str(name or "resume")).strip("_")
    return (cleaned or "resume")[:80]


def generate_resume_files(
    ats_check,
    rewrite_map=None,
    *,
    warnings=None,
    target_role="",
    career_level="auto",
    job_description="",
    matched_keywords=None,
):
    """Build DOCX + PDF from an ATSResumeCheck and save them on the model.

    The generated sections are deterministic unless ``rewrite_map`` contains a
    wording change that passes the evidence guard in ``build_sections``.  The
    ATS score is never recalculated here.
    """
    warnings = list(warnings or [])
    meta = {}
    try:
        sections = build_sections(
            ats_check.source_resume,
            rewrite_map,
            target_role=target_role or getattr(ats_check, "target_role", ""),
            career_level=career_level or getattr(ats_check, "career_level", "auto"),
            job_description=job_description or getattr(
                ats_check, "job_description", ""
            ),
            matched_keywords=matched_keywords,
            warnings=warnings,
            meta=meta,
        )
        docx_bytes = build_docx(sections)
        pdf_bytes = build_pdf(sections)
    except ImportError as exc:
        raise ResumeGenerationError(
            "Resume generation requires python-docx and reportlab."
        ) from exc
    header = sections[0] if sections and sections[0].get("is_contact") else {}
    detected_name = str(header.get("name") or "").strip()
    if (
        not (ats_check.student_name or "").strip()
        and detected_name
        and detected_name.lower() not in {"your name", "resume"}
    ):
        ats_check.student_name = detected_name
    filename_name = (
        detected_name
        if detected_name
        and detected_name.lower() not in {"your name", "resume"}
        else ats_check.student_name
    )
    base = (
        f"ats_resume_{_safe_filename(filename_name)}"
        f"_{ats_check.public_id.hex[:8]}"
    )
    docx_name = f"{base}.docx"
    pdf_name = f"{base}.pdf"
    try:
        ats_check.generated_docx.save(docx_name, ContentFile(docx_bytes), save=False)
        ats_check.generated_pdf.save(pdf_name, ContentFile(pdf_bytes), save=False)
    except Exception as exc:
        raise ResumeGenerationError(str(exc)[:200]) from exc
    ats_check.rewrite_warnings = warnings
    ats_check.rewrite_source = "llm" if meta.get("applied_rewrites") else "deterministic"
    return ats_check